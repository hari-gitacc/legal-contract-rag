import os
import re
from typing import List
from pypdf import PdfReader
from docx import Document as DocxReader
from langchain_core.documents import Document

class DocumentParser:
    @staticmethod
    def parse_pdf(file_path: str, filename: str) -> List[Document]:
        reader = PdfReader(file_path)
        documents = []
        for idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            if text.strip():
                documents.append(
                    Document(
                        page_content=text,
                        metadata={
                            "document_name": filename,
                            "page": idx + 1,
                            "section": f"Page {idx + 1}",
                            "file_path": file_path
                        }
                    )
                )
        return documents

    @staticmethod
    def parse_docx(file_path: str, filename: str) -> List[Document]:
        doc = DocxReader(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        full_text = "\n\n".join(paragraphs)
        return [
            Document(
                page_content=full_text,
                metadata={
                    "document_name": filename,
                    "page": 1,
                    "section": "Main Body",
                    "file_path": file_path
                }
            )
        ]

    @staticmethod
    def parse_txt(file_path: str, filename: str) -> List[Document]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        return [
            Document(
                page_content=text,
                metadata={
                    "document_name": filename,
                    "page": 1,
                    "section": "Document Text",
                    "file_path": file_path
                }
            )
        ]

    @classmethod
    def load_file(cls, file_path: str, filename: str) -> List[Document]:
        ext = os.path.splitext(filename)[1].lower()
        if ext == ".pdf":
            return cls.parse_pdf(file_path, filename)
        elif ext in [".docx", ".doc"]:
            return cls.parse_docx(file_path, filename)
        elif ext == ".txt":
            return cls.parse_txt(file_path, filename)
        else:
            raise ValueError(f"Unsupported file format: {ext}")

    @staticmethod
    def chunk_documents(documents: List[Document], chunk_size: int = 500, chunk_overlap: int = 75) -> List[Document]:
        chunked_docs: List[Document] = []
        for doc in documents:
            text = doc.page_content
            # Split by double newlines or punctuation boundaries
            splits = re.split(r'(?<=\n\n)|(?<=[.?!]\s)', text)
            current_chunk = []
            current_length = 0

            for part in splits:
                part_len = len(part)
                if current_length + part_len > chunk_size and current_chunk:
                    chunk_text = "".join(current_chunk).strip()
                    if chunk_text:
                        chunk_id = f"{doc.metadata['document_name']}_p{doc.metadata['page']}_c{len(chunked_docs)}"
                        citation_id = f"{doc.metadata['document_name']} | {doc.metadata['section']} | Page {doc.metadata['page']}"
                        chunked_docs.append(
                            Document(
                                page_content=chunk_text,
                                metadata={
                                    **doc.metadata,
                                    "chunk_id": chunk_id,
                                    "citation_id": citation_id
                                }
                            )
                        )
                    # Handle overlap
                    overlap_parts = []
                    overlap_len = 0
                    for prev_part in reversed(current_chunk):
                        if overlap_len + len(prev_part) <= chunk_overlap:
                            overlap_parts.insert(0, prev_part)
                            overlap_len += len(prev_part)
                        else:
                            break
                    current_chunk = overlap_parts
                    current_length = overlap_len

                current_chunk.append(part)
                current_length += part_len

            if current_chunk:
                chunk_text = "".join(current_chunk).strip()
                if chunk_text:
                    chunk_id = f"{doc.metadata['document_name']}_p{doc.metadata['page']}_c{len(chunked_docs)}"
                    citation_id = f"{doc.metadata['document_name']} | {doc.metadata['section']} | Page {doc.metadata['page']}"
                    chunked_docs.append(
                        Document(
                            page_content=chunk_text,
                            metadata={
                                **doc.metadata,
                                "chunk_id": chunk_id,
                                "citation_id": citation_id
                            }
                        )
                    )
        if chunked_docs:
            print(f"👉 [Parser] Created {len(chunked_docs)} chunks from {len(documents)} documents. First chunk preview: {chunked_docs[0].page_content[:100]!r}...")
        else:
            print(f"👉 [Parser] No chunks created from {len(documents)} documents.")
        return chunked_docs
